"""Validación humana a ciegas del ground truth (observación O1 de la revisión del 24-oct-2026).

  crear <pool.csv> <juicios.txt> <gt.csv> <dir_salida> [fraccion=0.10] [semilla=2026]
      Muestra estratificada por consulta (al menos el 10 % de los candidatos del pool que
      están en el corpus de la app, mínimo 5 por consulta) y escribe:
        Evaluador_A.xlsx y Evaluador_B.xlsx  (sin la decisión del modelo, en distinto orden)
        clave_muestra.csv                    (id → consulta, título, decisión del modelo; NO
                                              entregar: moverla a datos/ground_truth/)
  kappa <planilla_A.xlsx> <planilla_B.xlsx> <clave_muestra.csv>
      Lee las dos planillas completadas y calcula el kappa de Cohen entre los evaluadores y
      de cada uno con el modelo, en total y por consulta.
  aplicar <planilla_A.xlsx> <planilla_B.xlsx> <clave_muestra.csv> <gt.csv> <gt_salida.csv>
      Corrige el ground truth en las filas de la muestra por mayoría entre A, B y el modelo:
      cuando A y B coinciden y el modelo no, prevalece el juicio de los evaluadores.
"""
import csv
import random
import re
import sys
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


def leer_psv(ruta):
    with open(ruta, encoding="utf-8") as f:
        return list(csv.DictReader(f, delimiter="|", quoting=csv.QUOTE_NONE))


def criterios(ruta):
    c = {}
    for linea in open(ruta, encoding="utf-8"):
        m = re.match(r"# Q\d+ (.+?) :: (.+)", linea)
        if m:
            c[m.group(1).strip()] = m.group(2).strip()
    return c


def norm(t):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9áéíóúüñ ]", " ", t.lower())).strip()


INSTRUCCIONES = [
    "Validación del ground truth de la evaluación de APPPapers",
    "",
    "1. Cada fila es una publicación candidata para una consulta. Lea el criterio de la consulta, el título, "
    "las palabras clave y el resumen.",
    "2. En la columna «¿Relevante?» elija Sí si la publicación cumple el criterio y No si no lo cumple.",
    "3. Juzgue cada fila por separado, sin consultar a la otra persona evaluadora ni otras fuentes. "
    "No hay una cantidad esperada de Sí o de No.",
    "4. Si duda, decida igual y explique la duda en «Comentario».",
    "5. No modifique las demás columnas. Al terminar, guarde el archivo y devuélvalo.",
]


def escribir_planilla(ruta, filas, evaluador):
    wb = Workbook()
    ins = wb.active
    ins.title = "Instrucciones"
    for i, t in enumerate(INSTRUCCIONES + ["", f"Evaluador/a: {evaluador}", "Nombre:", "Formación:"], 1):
        ins.cell(i, 1, t).alignment = Alignment(wrap_text=True)
    ins["A1"].font = Font(bold=True, size=13)
    ins.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Juicios")
    enc = ["N°", "Consulta", "Criterio de relevancia", "Título", "Año", "Palabras clave", "Resumen",
           "¿Relevante?", "Comentario"]
    ws.append(enc)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for f in filas:
        ws.append([f["id"], f["query"], f["criterio"], f["titulo"], f["ano"], f["palabras_clave"],
                   f["resumen"], None, None])
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[7].fill = amarillo
        row[8].fill = amarillo
    for col, w in zip("ABCDEFGHI", [6, 22, 40, 45, 6, 30, 80, 12, 30]):
        ws.column_dimensions[col].width = w
    dv = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    dv.add(f"H2:H{len(filas) + 1}")
    ws.add_data_validation(dv)
    ws.freeze_panes = "B2"
    wb.save(ruta)


def crear(pool_csv, juicios_txt, gt_csv, salida, fraccion=0.10, semilla=2026):
    import os
    os.makedirs(salida, exist_ok=True)
    pool = [p for p in leer_psv(pool_csv) if p["en_corpus_app"] == "TRUE"]
    gt = {g["query"]: {norm(t) for t in g["relevantes_titulos"].split(";;")} for g in leer_psv(gt_csv)}
    crit = criterios(juicios_txt)
    rnd = random.Random(semilla)
    por_q = defaultdict(list)
    for p in pool:
        por_q[p["query"]].append(p)
    muestra = []
    for q, cands in por_q.items():
        k = min(len(cands), max(5, round(fraccion * len(cands) + 0.4999)))
        muestra += rnd.sample(cands, k)
    print(f"Pool: {len(pool)} candidatos | muestra: {len(muestra)} ({100 * len(muestra) / len(pool):.1f} %)")
    rnd.shuffle(muestra)
    filas = []
    for i, p in enumerate(muestra, 1):
        filas.append({"id": i, "query": p["query"], "criterio": crit[p["query"]], "titulo": p["titulo"],
                      "ano": p["ano"], "palabras_clave": p["palabras_clave"][:400], "resumen": p["resumen"][:900],
                      "modelo": "Sí" if norm(p["titulo"]) in gt[p["query"]] else "No"})
    with open(f"{salida}/clave_muestra.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter="|")
        w.writerow(["id", "query", "titulo", "decision_modelo"])
        for fl in filas:
            w.writerow([fl["id"], fl["query"], fl["titulo"], fl["modelo"]])
    escribir_planilla(f"{salida}/Evaluador_A.xlsx", filas, "A")
    orden_b = filas[:]
    random.Random(semilla + 1).shuffle(orden_b)
    escribir_planilla(f"{salida}/Evaluador_B.xlsx", orden_b, "B")
    print("Sí del modelo en la muestra:", sum(f["modelo"] == "Sí" for f in filas))


def kappa(a, b):
    n = len(a)
    if n == 0:
        return float("nan")
    po = sum(x == y for x, y in zip(a, b)) / n
    pa, pb = sum(a) / n, sum(b) / n
    pe = pa * pb + (1 - pa) * (1 - pb)
    return 1.0 if pe == 1 else (po - pe) / (1 - pe)


def leer_juicios(ruta):
    ws = load_workbook(ruta)["Juicios"]
    r = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        v = (row[7] or "").strip().lower()
        if v not in ("sí", "si", "no"):
            raise ValueError(f"{ruta}: fila N° {row[0]} sin juicio")
        r[int(row[0])] = v != "no"
    return r


def calcular(ruta_a, ruta_b, ruta_clave):
    clave = leer_psv(ruta_clave)
    a, b = leer_juicios(ruta_a), leer_juicios(ruta_b)
    if len(set(a.values())) == 1 or len(set(b.values())) == 1:
        print("ADVERTENCIA: una planilla tiene todas las filas con el mismo juicio; revisar antes de usarla.")
    ids = [int(c["id"]) for c in clave]
    m = {int(c["id"]): c["decision_modelo"] == "Sí" for c in clave}

    def resumen(sub, etiqueta):
        A, B, M = [a[i] for i in sub], [b[i] for i in sub], [m[i] for i in sub]
        acu = lambda x, y: 100 * sum(p == q for p, q in zip(x, y)) / len(x)
        print(f"{etiqueta:52s} n={len(sub):3d} | A-B {acu(A, B):5.1f} % k={kappa(A, B):5.2f} | "
              f"A-modelo {acu(A, M):5.1f} % k={kappa(A, M):5.2f} | B-modelo {acu(B, M):5.1f} % k={kappa(B, M):5.2f}")

    resumen(ids, "TOTAL")
    for q in sorted({c["query"] for c in clave}):
        resumen([int(c["id"]) for c in clave if c["query"] == q], q)
    acuerdo = [i for i in ids if a[i] == b[i]]
    difiere = sum(a[i] != m[i] for i in acuerdo)
    print(f"\nFilas en que A y B coinciden: {len(acuerdo)}; de ellas, el modelo difiere en {difiere}.")
    print(f"Sí: A {sum(a.values())}, B {sum(b.values())}, modelo {sum(m.values())} de {len(ids)}.")


def aplicar(ruta_a, ruta_b, ruta_clave, gt_csv, salida):
    a, b = leer_juicios(ruta_a), leer_juicios(ruta_b)
    with open(gt_csv, encoding="utf-8") as f:
        cab = f.readline().rstrip("\n").split("|")
    gt = leer_psv(gt_csv)
    por_q = {g["query"]: g for g in gt}
    for c in leer_psv(ruta_clave):
        i, q, titulo = int(c["id"]), c["query"], c["titulo"]
        modelo = c["decision_modelo"] == "Sí"
        if a[i] != b[i] or a[i] == modelo:
            continue
        titulos = [t for t in por_q[q]["relevantes_titulos"].split(";;") if t]
        if a[i]:
            titulos.append(titulo)
        else:
            titulos = [t for t in titulos if norm(t) != norm(titulo)]
        por_q[q]["relevantes_titulos"] = ";;".join(titulos)
        por_q[q]["n_relevantes"] = str(len(titulos))
        print(f"N° {i:3d} → {'Sí' if a[i] else 'No'} | {q} | {titulo[:80]}")
    with open(salida, "w", encoding="utf-8", newline="") as f:
        f.write("|".join(cab) + "\n")
        for g in gt:
            f.write("|".join(g[k] for k in cab) + "\n")
    total = sum(int(g["n_relevantes"]) for g in gt)
    print(f"Ground truth corregido: {total} relevantes → {salida}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    if sys.argv[1] == "crear":
        crear(*sys.argv[2:6], *[float(x) for x in sys.argv[6:7]], *[int(x) for x in sys.argv[7:8]])
    elif sys.argv[1] == "kappa":
        calcular(*sys.argv[2:5])
    else:
        aplicar(*sys.argv[2:7])
